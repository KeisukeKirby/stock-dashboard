"""Excel に直接入力された Return (返品) をダッシュボードに取り込むための JSON を作る。

使い方:
    python scripts/import_returns.py data/VFF_Stock_with_returns_20260929_1620.xlsx
      -> data/returns_import.json  (convert.py の 7 つ目の引数で stock.json の returnsSeed に入る)

ダッシュボードは初回表示時に、この返品を「返品輸送中」として共有データに 1 度だけ登録する
(同じマスに既に入力があればそちらを優先)。オフィス在庫へは「受領完了」を押すまで加算しない。

黄色 (FFFF00) で塗られたマスは highlights として記録し、ダッシュボードでも同じマスを黄色で表示する。
"""
import hashlib
import json
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "data" / "returns_import.json"


def main(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    h3 = [c.value for c in ws[3]]
    h4 = [c.value for c in ws[4]]
    cols, kinds, grp = [], {}, None
    for i, (a, b) in enumerate(zip(h3, h4)):
        grp = a or grp
        if b in ("Quantity", "Return"):
            kinds[i] = (grp, "qty" if b == "Quantity" else "ret")
        if b == "Return" and grp not in ("Store Total", "Company Total", "Total"):
            cols.append((i, grp))
    returns, highlights = {}, []
    for row in ws.iter_rows(min_row=5):
        r = [c.value for c in row]
        if not r[0] or r[0] == "Total":
            continue
        code = str(r[0]).strip()
        for i, store in cols:
            q = r[i]
            if q not in (None, "", 0):
                returns[f"{code}@{store}"] = int(q)
        for i, c in enumerate(row):
            if i in kinds and c.fill and c.fill.fill_type == "solid" and c.fill.fgColor.rgb == "FFFFFF00":
                store, kind = kinds[i]
                highlights.append({"code": code, "store": store, "kind": kind})
    body = json.dumps(returns, sort_keys=True, ensure_ascii=False)
    out = {
        "id": f"{Path(path).stem}-{hashlib.sha1(body.encode()).hexdigest()[:8]}",
        "source": Path(path).name,
        "returns": returns,
        "highlights": highlights,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(returns)} cells, {sum(returns.values())} pairs, {len(highlights)} highlighted cells -> {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
