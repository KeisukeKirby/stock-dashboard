"""イベント (Event Asok) の残り在庫を計算する。

  残り在庫 = スタート在庫 (EVENT_*.xlsx の「Event Asoke」列) − 販売数 (POS 注文明細、締め日の営業終了まで)

使い方:
    python scripts/event_stock.py data/EVENT_SeP2026_start_stock.xlsx 2026-09-25 data/event_orders/*.xlsx

結果は data/event_asok_stock.json に保存され、convert.py の 5 つ目の引数に渡すと
Event Asok 列に加算される (Office は動かさない)。

注文明細の扱い:
  - 1 注文が複数行にまたがるので、注文番号・日付・状態・倉庫は上の行から引き継ぐ
  - 取消 (Voided) は除外、倉庫が「Event 1」以外 (Online・Paradise Park など) は除外
    (倉庫列のないファイルはイベント POS のみの出力として扱う)
  - 締め日より後の日付は除外
  ※ 注文明細は顧客情報を含むため data/event_orders/ は Git に入れない
"""
import collections
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "data" / "event_asok_stock.json"
SKU_RE = re.compile(r"^VFF\d{4}\(.+\)$")


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, str) and v.strip():
        d, m, y = v.split()[0].split("/")
        return datetime.date(int(y), int(m), int(d))
    return None


def read_start(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "").replace("\n", " ").strip() for h in rows[0]]
    c_qty = next(i for i, h in enumerate(head) if "event" in h.lower())
    start = collections.Counter()
    for r in rows[1:]:
        if not r[0]:
            continue
        code = str(r[0]).strip()
        if code.count("(") > code.count(")"):  # Excel 上で ")" が欠けているコードを補う
            code += ")"
        start[code] += int(r[c_qty] or 0)
    return start


def read_sales(files, cutoff):
    sold = collections.Counter()
    for f in files:
        ws = openpyxl.load_workbook(f, data_only=True, read_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        head = [str(h).strip() if h else "" for h in rows[1]]
        col = lambda n: head.index(n) if n in head else None
        c_order, c_date, c_status, c_wh = col("Sales order No."), col("Date"), col("Status"), col("Warehouse/Branch")
        c_code, c_qty = col("Product code"), col("Quantity")
        cur = {}
        for r in rows[2:]:
            if r[c_order]:
                cur = {"date": parse_date(r[c_date]),
                       "status": r[c_status] if c_status is not None else None,
                       "wh": r[c_wh] if c_wh is not None else None}
            code = r[c_code]
            if not code or cur.get("status") == "Voided" or cur.get("wh") not in (None, "Event 1"):
                continue
            if cur["date"] is None or cur["date"] > cutoff:
                continue
            sold[str(code).strip()] += int(float(r[c_qty] or 0))
    return sold


def main(start_path, cutoff, *order_files):
    cutoff = datetime.date.fromisoformat(cutoff)
    start = read_start(start_path)
    sold = read_sales(order_files, cutoff)
    remain, oversold, not_in_start = {}, {}, {}
    for code, q in start.items():
        s = sold.get(code, 0)
        remain[code] = max(q - s, 0)
        if s > q:
            oversold[code] = {"start": q, "sold": s}
    for code, s in sold.items():
        if SKU_RE.match(code) and code not in start:
            not_in_start[code] = s
    out = {
        "cutoff": str(cutoff),
        "startSource": Path(start_path).name,
        "orderSources": [Path(f).name for f in order_files],
        "startTotal": sum(start.values()),
        "soldFromStart": sum(min(sold.get(c, 0), q) for c, q in start.items()),
        "remainTotal": sum(remain.values()),
        "remain": remain,
        "oversold": oversold,            # スタート在庫より多く売れた (0 として扱った)
        "soldNotInStart": not_in_start,  # スタート在庫にない商品の販売 (引けないので対象外)
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"start {out['startTotal']} - sold {out['soldFromStart']} = remain {out['remainTotal']} -> {OUT}")
    print(f"oversold: {oversold}")
    print(f"sold but not in start stock: {sum(not_in_start.values())} pairs / {len(not_in_start)} SKUs")


if __name__ == "__main__":
    main(*sys.argv[1:])
