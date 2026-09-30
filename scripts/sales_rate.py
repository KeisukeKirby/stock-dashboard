"""店舗別・SKU 別の月平均販売足数 (VFF シューズ) を計算して data/sales_rate.json に保存する。

在庫ダッシュボードの「月平均販売」「在庫月数」列に使う (convert.py の 6 つ目の引数)。

使い方:
    python scripts/sales_rate.py            # data/sales_raw/ の下記ファイルを読む

元データ (data/sales_raw/、顧客情報を含むため Git 管理外):
  K Village / Central LP : order_detail_*_7duk.xlsx  (EDV の受注明細。倉庫 Kvillage / Coollabo Cen LP 3F)
  Paradise Park          : order_detail_*_dint.xlsx  (Barefoot の受注明細。倉庫 Paradise Park)
  Central CL (Chidlom)   : BFT_Central_Total_Department_Jan-Jun_26_new.xlsx (Export シート)
                           + BFT_Central_Total_Department_07-2026.xlsx (Export シート, 7 月)
                           + BFT_Sale_Online_Shopee_Lazada_Paradise_Central_Jul-Aug.xlsx (Central シート, 8 月)
  Online (Office 列)     : order_detail_*_dint.xlsx のオンライン注文 (Shopee / Lazada / Facebook / LINE /
                           Instagram / Website、販売チャネル空欄の TX 注文、倉庫 Online)
  Event (Company Total)  : 両受注明細の倉庫 Event 1 / Event 2 + CART Central LP のイベント分
                           (R# RC-127- のレシート、5/21-5/24 の RV 注文)
  Siam Discovery         : Sales_Siam_Dis_Jan-Jun_26.xlsx + BFT_Siam_Discovery_Jul-Aug_26.xlsx

集計ルール:
  - VFF シューズのみ (サイズが数字の商品。靴下・Furoshiki 等は除く)、取消 (Voided) は除外
  - 店舗の列は店頭の販売のみ。店舗在庫から発送したオンライン注文は「Online」に数える
    (EDV 販売ダッシュボードの「K village」と同じ定義)
  - Online / Event は受注明細の全期間 (1 月〜) の月数で割る (販売のない月も 0 として数える)
  - 上記以外 (本社倉庫のチャネル空欄の販売、EDV の CART Central LP の通常販売、委託・輸出) は含めない
  - Central 百貨店は CHIDLOM のみ (CHIDLOM ONLINE・Central World・Lardprao・Eastville は含めない)
  - 月平均 = 期間の販売足数 ÷ データのある月数
"""
import collections
import datetime
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "sales_raw"
OUT = ROOT / "data" / "sales_rate.json"
STOCK = ROOT / "public" / "stock.json"

F_EDV = "order_detail_202609291711_7duk.xlsx"
F_BFT = "order_detail_202609301006_dint.xlsx"
F_CEN_H1 = "BFT_Central_Total_Department_Jan-Jun_26_new.xlsx"
F_CEN_JUL = "BFT_Central_Total_Department_07-2026.xlsx"
F_CEN_H2 = "BFT_Sale_Online_Shopee_Lazada_Paradise_Central_Jul-Aug.xlsx"
F_SIAM_H1 = "Sales_Siam_Dis_Jan-Jun_26.xlsx"
F_SIAM_H2 = "BFT_Siam_Discovery_Jul-Aug_26.xlsx"

KV, LP, PP, CL, SD = "K Village", "Central LP", "Paradise Park", "Central CL (Department)", "Siam Discovery"
ONLINE, EVENT = "Online", "Event"  # 全期間で割る区分
STORE_WAREHOUSE = {"Kvillage": KV, "Coollabo Cen LP 3F": LP, "Paradise Park": PP}
ONLINE_CHANNELS = {"Shopee", "Shopee VFF", "Shopee BFI", "Lazada", "Facebook", "LINE", "Instagram", "Website", "barefootinc"}
SHOE_SIZE = re.compile(r"^[MWU]?\d+(?:\.\d+)?$")
CODE_RE = re.compile(r"^VFF(\d{4})\((.+),([MWU]?\d+)\)$")
N = lambda s: re.sub(r"[\s/.\-]", "", str(s).upper()).replace("DARKGY", "DARKGRAY")

# 百貨店・Siam の売上表の表記 -> 在庫表の表記
MODEL_ALIAS = {"SPIDWALK": "SPIDRWALK", "TRAILPOE": "TRAILOPE", "VTRAIN": "VTRAIN20"}
ABBR_ALIAS = {  # (モデル番号, 売上表の略号) -> 在庫表の略号
    (17, "MIL"): "MILDGY", (5, "LIG"): "LIGGY", (21, "FU"): "IVFU",
    (26, "BK"): "TTBK", (2, "BKFI"): "BKF", (23, "LIGNBK"): "LGNBK",
}


def vff_shoe(name):
    """'VFF V-Soul(W37, Nude)' -> (size, color)。靴以外は None。"""
    m = re.search(r"\(([^()]*)\)\s*$", str(name or ""))
    if not m:
        return None
    parts = [p.strip() for p in m.group(1).split(",")]
    return (parts[0], ", ".join(parts[1:])) if SHOE_SIZE.match(parts[0]) else None


def parse_dmy(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    d, m, y = (int(x) for x in str(v).split()[0].split("/"))
    return datetime.date(y, m, d)


def read_orders(path, sales, meta, periods):
    """受注明細 (order_detail_*.xlsx)。VFF シューズを 店舗 / Online / Event・月・コード別に数える。"""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[1]) if h}
    g = lambda r, k: r[ix[k]]
    for r in rows[2:]:
        if g(r, "Type") != "Sell" or g(r, "Status") == "Voided":
            continue
        code, name = g(r, "Product code"), g(r, "Product name")
        if not code or not str(code).startswith("VFF") or not vff_shoe(name):
            continue
        wh = g(r, "Warehouse/Branch") or ""
        ch, order_no = g(r, "Sales channel") or "", str(g(r, "Sales order No.") or "")
        date = parse_dmy(g(r, "Date"))
        if ch in ONLINE_CHANNELS or (not ch and order_no.startswith("TX")) or wh == "Online":
            store = ONLINE
        elif wh in ("Event 1", "Event 2") or (wh == "CART Central LP" and (
                order_no.startswith("R# RC-127-")
                or (order_no.startswith("RV") and datetime.date(2026, 5, 21) <= date <= datetime.date(2026, 5, 24)))):
            store = EVENT
        else:
            store = STORE_WAREHOUSE.get(wh)
        month = date.strftime("%Y-%m")
        periods.add(month)
        if not store:
            continue
        code = str(code).strip()
        sales[code][(store, month)] += float(g(r, "Quantity") or 0)
        meta.setdefault(code, name)


class SkuMatcher:
    """百貨店・Siam の売上表の商品表記を在庫表の商品コードに合わせる。"""

    def __init__(self, known_codes):
        self.by_abbr, self.by_name, self.model_no = {}, {}, {}
        for code, (model, color) in known_codes.items():
            m = CODE_RE.match(code)
            if not m:
                continue
            p, ab, size = int(m.group(1)), m.group(2), m.group(3)
            self.by_abbr.setdefault((p, N(ab), size), code)
            mdl = N(model.replace("VFF ", ""))
            self.by_name.setdefault((mdl, N(color), size), code)
            self.model_no.setdefault(mdl, p)

    def abbr(self, p, ab, size):
        return self.by_abbr.get((p, ABBR_ALIAS.get((p, N(ab)), N(ab)), size))

    def central(self, code, name):
        m = re.match(r"^VFF(\d+)\(([^,]+),([MWU]?\d+)\)$", str(code).strip())
        if m:
            hit = self.abbr(int(m.group(1)), m.group(2), m.group(3))
            if hit:
                return hit
        n = re.match(r"^(.*?)\(\s*([MWU]?\d+)\s*,\s*(.+)\)\s*$", str(name))
        if n:
            mdl = MODEL_ALIAS.get(N(n.group(1)), N(n.group(1)))
            return self.by_name.get((mdl, N(n.group(3)), n.group(2)))
        return None

    def siam(self, item):
        m = re.match(r"^(.*?)\s*\(\s*([^,]+?)\s*,\s*([MWU]?\d+)\s*\)", str(item))  # 末尾の注記は無視
        if not m:
            return None
        mdl = MODEL_ALIAS.get(N(m.group(1)), N(m.group(1)))
        p = self.model_no.get(mdl)
        return self.abbr(p, m.group(2), m.group(3)) if p is not None else None


def is_shoe_text(name):
    s = str(name).upper()
    return bool(re.search(r"[\(,]\s*[MWU]?\d{2}\s*[,)]", s)) and not re.search(r"SOCK|OLENO|OLN|BFJ|TABI", s)


MON = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}


def read_central(sales, matcher, unmatched):
    for fname, sheet, (c_store, c_name, c_code, c_mon, c_qty) in [
        (F_CEN_H1, "Export", (0, 1, 2, 3, 4)),
        (F_CEN_JUL, "Export", (0, 1, 2, 3, 4)),
        (F_CEN_H2, "Central", (0, 3, 2, 1, 4)),
    ]:
        rows = list(openpyxl.load_workbook(RAW / fname, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
        for r in rows[1:]:
            if r[c_store] != "CHIDLOM" or not str(r[c_code] or "").startswith("VFF") or not is_shoe_text(r[c_name]):
                continue
            mon, year = str(r[c_mon]).split("-")
            month = f"{year}-{MON[mon[:3]]:02d}"
            code = matcher.central(r[c_code], r[c_name])
            if code:
                sales[code][(CL, month)] += float(r[c_qty] or 0)
            else:
                unmatched.append([CL, month, str(r[c_code]), str(r[c_name]), r[c_qty]])


def read_siam(sales, matcher, unmatched):
    for fname in (F_SIAM_H1, F_SIAM_H2):
        rows = list(openpyxl.load_workbook(RAW / fname, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
        for r in rows[1:]:
            d, item, q = r[0], r[1], r[2]
            if not item or not is_shoe_text(item):
                continue
            month = d.strftime("%Y-%m")
            code = matcher.siam(item)
            if code:
                sales[code][(SD, month)] += float(q or 0)
            else:
                unmatched.append([SD, month, "", str(item), q])


def main():
    sales = collections.defaultdict(collections.Counter)
    meta, periods = {}, set()
    read_orders(RAW / F_EDV, sales, meta, periods)
    read_orders(RAW / F_BFT, sales, meta, periods)

    stock = json.loads(STOCK.read_text(encoding="utf-8"))
    known = {i["code"]: (i["model"], i["color"]) for i in stock["items"] if i["code"]}
    for code, name in meta.items():  # 受注明細にだけある商品も照合対象に
        if code not in known:
            known[code] = (str(name).split("(")[0].strip(), vff_shoe(name)[1])
    matcher = SkuMatcher(known)
    unmatched = []
    read_central(sales, matcher, unmatched)
    read_siam(sales, matcher, unmatched)

    months_present = collections.defaultdict(set)
    for c in sales.values():
        for (s, m), q in c.items():
            if q:
                months_present[s].add(m)
    months = {s: sorted(v) for s, v in months_present.items()}
    for s in (ONLINE, EVENT):
        months[s] = sorted(periods)
    rate = {}
    for code, c in sales.items():
        per = {}
        for s, ms in months.items():
            total = sum(c.get((s, m), 0) for m in ms)
            if total:
                per[s] = round(total / len(ms), 3)
        if per:
            rate[code] = per
    totals = {s: sum(q for c in sales.values() for (st, _), q in c.items() if st == s) for s in months}
    out = {
        "generatedAt": datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=7))).strftime("%Y-%m-%d %H:%M"),
        "months": months,
        "totals": totals,
        "rate": rate,
        "unmatched": unmatched,
        "sources": [F_EDV, F_BFT, F_CEN_H1, F_CEN_JUL, F_CEN_H2, F_SIAM_H1, F_SIAM_H2],
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    for s in months:
        print(f"{s:24s} {totals[s]:6.0f} pairs / {len(months[s])} months ({months[s][0]}..{months[s][-1]})")
    print(f"{len(rate)} SKUs -> {OUT}; unmatched {len(unmatched)}: {unmatched}")


if __name__ == "__main__":
    main()
