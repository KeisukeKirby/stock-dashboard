"""Excel (Store_Stock_*.xlsx) の「VFF Shoes Stock」シートを public/stock.json に変換する。
オフィス在庫の Excel (VFF_Stock_*.xlsx) を渡すと「Office」列として追加する。
新入荷の店舗配分表 (「配分表」シートのある Excel) を渡すと、配分した足数を Office から各店舗へ移す。

使い方:
    pip install openpyxl
    python scripts/convert.py data/Store_Stock_092526.xlsx data/VFF_Stock_25-09-26.xlsx \
        [data/New_Arrival_Allocation.xlsx] [data/import_9.26_move_to_branch.xlsx]

4 つ目に店舗移動表 (import_*_move_to_branch.xlsx) を渡すと、その「Asok」列の数を
「Event Asok」列として追加し、Office から移す。空欄の Code も移動表のコードで補完する。

Excel 側で保存時に計算済みの値 (data_only) を読み込むため、
Excel で一度保存したファイルを使ってください。
"""
import json
from datetime import datetime, timedelta, timezone
import re
import shutil
import sys
from pathlib import Path

import openpyxl

SHEET = "VFF Shoes Stock"
PUBLIC = Path(__file__).resolve().parent.parent / "public"
OUT = PUBLIC / "stock.json"
XLSX_COPY = PUBLIC / "source.xlsx"  # 画面の「元の Excel をダウンロード」用
OFFICE_XLSX_COPY = PUBLIC / "source_office.xlsx"

# オフィス在庫シートの列 (1 始まり)
OFFICE_COL_ITEM = 2    # B: ITMES  例) "V-Run M\n(Black/Yellow)" (縦に結合)
OFFICE_COL_SIZE = 3    # C: Size
OFFICE_COL_STOCK = 66  # BN: Stock > Office
OFFICE_NAME = "Office"


def num(v):
    return int(v) if isinstance(v, (int, float)) else 0


def cell(v):
    """空セルは None (表示も空欄) のまま残す。"""
    return int(v) if isinstance(v, (int, float)) else None


def norm_color(s):
    """表記ゆれ (空白・ピリオド・大文字小文字・Balck) を吸収して比較する。"""
    return re.sub(r"[\s.]", "", s.lower().replace("balck", "black"))


def read_office(path):
    """オフィス在庫を [(model, color, size, qty)] で返す (表示順のまま)。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    out = []
    item = None
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, OFFICE_COL_ITEM).value
        size = ws.cell(r, OFFICE_COL_SIZE).value
        if name:
            item = str(name)
        if item is None or size is None or str(size).strip().upper() == "TOTAL" or "\n" not in item:
            continue
        head, color = item.split("\n", 1)
        model = "VFF " + re.sub(r"\s+[MWU]$", "", head.strip())  # "V-Run M" -> "VFF V-Run"
        color = color.strip().strip("()").strip().replace("Balck", "Black")
        size = str(size).strip()
        if head.strip().endswith(" U") and size.isdigit():  # Scramkey U: 40 -> U40
            size = "U" + size
        out.append((model, color, size, cell(ws.cell(r, OFFICE_COL_STOCK).value)))
    return out


CODE_RE = re.compile(r"^(VFF\d+)\((.+),([^,()]+)\)$")  # 例) VFF0002(BK/YL,M42)


def guess_code(items, model, color, size):
    """同じモデル・カラーの商品コードがあれば、サイズ部分だけ差し替えてコードを作る。
    (モデル番号とカラー略号が確実に分かる場合のみ。分からなければ空欄)"""
    for x in items:
        if x["code"] and x["model"] == model and norm_color(x["color"]) == norm_color(color):
            m = CODE_RE.match(x["code"])
            if m:
                return f"{m.group(1)}({m.group(2)},{size})"
    return ""


def add_office(items, office):
    """店舗の行に Office 在庫を追加。店舗にない商品は、在庫がある場合のみ行を追加する。"""
    index = {(i["model"], norm_color(i["color"]), i["size"]): i for i in items}
    for i in items:
        i["qty"].append(0)
        i["ret"].append(None)
    matched, added = 0, 0
    for model, color, size, qty in office:
        it = index.get((model, norm_color(color), size))
        if it:
            it["qty"][-1] = qty or 0
            matched += 1
        elif qty:
            new = {"code": guess_code(items, model, color, size), "model": model, "color": color, "size": size,
                   "qty": [0] * (len(items[0]["qty"]) - 1) + [qty],
                   "ret": [None] * len(items[0]["ret"]), "office_only": True}
            # 同じモデル・カラーの行の後ろ (なければ同じモデルの後ろ、それもなければ末尾) に入れる
            same = [n for n, x in enumerate(items) if x["model"] == model]
            same_color = [n for n in same if norm_color(items[n]["color"]) == norm_color(color)]
            pos = (same_color or same or [len(items) - 1])[-1] + 1
            items.insert(pos, new)
            index[(model, norm_color(color), size)] = new
            added += 1
    for i in items:
        i["totalQty"] = sum(i["qty"])
        i["totalRet"] = sum(v or 0 for v in i["ret"])
        i.pop("office_only", None)
    return matched, added


ALLOC_SHEET = "配分表"


def store_match(name, stores):
    """配分表の店舗名 (例: "Central CL (Chidlom)") をダッシュボードの店舗名に合わせる。"""
    base = re.sub(r"\s*\(.*\)$", "", str(name)).strip().lower()
    for s in stores:
        if re.sub(r"\s*\(.*\)$", "", s).strip().lower() == base:
            return s
    raise SystemExit(f"配分表の店舗名が一致しません: {name}")


def apply_allocation(items, store_names, path):
    """配分表の店舗配分を反映する。新入荷はオフィス在庫に含まれているので、
    配分した足数を Office から各店舗へ移す (Company Total は変わらない)。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[ALLOC_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(n for n, r in enumerate(rows) if r and r[0] == "モデル" and "入荷数" in r)
    head = rows[hi]
    c_model, c_color, c_size = head.index("モデル"), head.index("カラー"), head.index("サイズ")
    c_total = head.index("店舗合計")
    c_stores = list(range(head.index("配分グループ") + 1, c_total))
    col_to_idx = {c: store_names.index(store_match(head[c], store_names)) for c in c_stores}
    office = store_names.index(OFFICE_NAME)
    index = {(i["model"], norm_color(i["color"]), i["size"]): i for i in items}
    moved, n = 0, 0
    for r in rows[hi + 1:]:
        if not r or not r[c_model] or r[c_model] == "合計" or not isinstance(r[c_total], (int, float)):
            continue
        it = index.get((r[c_model], norm_color(str(r[c_color])), str(r[c_size]).strip()))
        if not it:
            raise SystemExit(f"配分表の商品がダッシュボードにありません: {r[c_model]} {r[c_color]} {r[c_size]}")
        total = 0
        for c, k in col_to_idx.items():
            q = int(r[c] or 0)
            it["qty"][k] += q
            total += q
        if total != int(r[c_total]):
            raise SystemExit(f"店舗合計が一致しません: {r[c_model]} {r[c_color]} {r[c_size]}")
        if it["qty"][office] < total:
            raise SystemExit(f"Office 在庫が配分数より少ない: {r[c_model]} {r[c_color]} {r[c_size]}")
        it["qty"][office] -= total
        moved += total
        n += 1
    for i in items:
        i["totalQty"] = sum(i["qty"])
    print(f"allocation: {n} rows, {moved} pairs moved from Office to stores")
    return moved


EVENT_NAME = "Event Asok"
EVENT_COL = "Asok"


def apply_event(items, store_names, path):
    """店舗移動表の Asok 列を Event Asok 列として追加し、Office から移す。
    商品は「ชื่อสินค้า」(例: VFF Groundsplay LS(M40, Black/Lime)) またはコードで照合する。"""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    c_code, c_name, c_event = 0, 1, head.index(EVENT_COL)
    office = store_names.index(OFFICE_NAME)
    store_names.append(EVENT_NAME)
    for i in items:
        i["qty"].append(0)
        i["ret"].append(None)
    by_code = {i["code"]: i for i in items if i["code"]}
    norm = lambda c: re.sub(r"[\s./]", "", c.lower())
    by_key = {(i["model"], norm(i["color"]), i["size"]): i for i in items}
    moved, codes = 0, 0
    for r in rows[1:]:
        if not r or not r[c_code]:
            continue
        m = re.match(r"^(.*)\((\w+), (.+)\)$", str(r[c_name]).strip())
        if not m:
            raise SystemExit(f"移動表の商品名を読めません: {r[c_name]}")
        model, size, color = m.group(1).strip(), m.group(2), m.group(3).strip()
        it = by_code.get(str(r[c_code]).strip()) or by_key.get((model, norm(color), size))
        if not it:
            raise SystemExit(f"移動表の商品がダッシュボードにありません: {r[c_name]}")
        if not it["code"]:
            it["code"] = str(r[c_code]).strip()
            codes += 1
        q = int(r[c_event] or 0)
        if it["qty"][office] < q:
            raise SystemExit(f"Office 在庫が移動数より少ない: {r[c_name]}")
        it["qty"][-1] += q
        it["qty"][office] -= q
        moved += q
    for i in items:
        i["totalQty"] = sum(i["qty"])
    print(f"event: {moved} pairs moved from Office to {EVENT_NAME}, {codes} codes filled")
    return moved


def apply_event_stock(items, store_names, path):
    """scripts/event_stock.py の結果 (スタート在庫 − 販売) を Event Asok 列に加える。Office は動かさない。"""
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    k = store_names.index(EVENT_NAME)
    by_code = {i["code"]: i for i in items if i["code"]}
    added, missing = 0, []
    for code, q in data["remain"].items():
        it = by_code.get(code)
        if not it:
            missing.append(code)
            continue
        it["qty"][k] += q
        added += q
    if missing:
        raise SystemExit(f"イベント在庫の商品がダッシュボードにありません: {missing}")
    for i in items:
        i["totalQty"] = sum(i["qty"])
    print(f"event stock: +{added} pairs to {EVENT_NAME} (cutoff {data['cutoff']})")
    return data, added


def main(path, office_path=None, alloc_path=None, event_path=None, event_stock_path=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))

    title = rows[0][0] or "Inventory"
    as_of = rows[1][0] or ""
    header = rows[2]

    # 店舗は E 列から 2 列ずつ (Quantity / Return)。"Total" 列の手前まで。
    stores = []
    col = 4
    while col < len(header) and header[col] and str(header[col]).strip() != "Total":
        stores.append({"name": str(header[col]).strip(), "col": col})
        col += 2

    total_col = col  # "Total" 列 (Quantity, Return)
    items = []
    total_row = None
    for r in rows[4:]:
        code = r[0]
        if not code:
            continue
        if str(code).strip() == "Total":
            total_row = {
                "qty": [cell(r[s["col"]]) for s in stores] + [cell(r[total_col])],
                "ret": [cell(r[s["col"] + 1]) for s in stores] + [cell(r[total_col + 1])],
            }
            continue
        qty = [num(r[s["col"]]) for s in stores]
        ret = [cell(r[s["col"] + 1]) for s in stores]
        items.append({
            "code": str(code).strip(),
            "model": (r[1] or "").strip(),
            "color": (r[2] or "").strip(),
            "size": str(r[3] or "").strip(),
            "qty": qty,
            "ret": ret,
            "totalQty": cell(r[total_col]),
            "totalRet": cell(r[total_col + 1]),
        })

    store_names = [s["name"] for s in stores]
    office_src = None
    if office_path:
        matched, added = add_office(items, read_office(office_path))
        store_names.append(OFFICE_NAME)
        # Office を一番左 (先頭) に並べ替える
        store_names.insert(0, store_names.pop())
        for i in items:
            i["qty"].insert(0, i["qty"].pop())
            i["ret"].insert(0, i["ret"].pop())
        n = len(store_names)
        alloc_src, alloc_moved = None, 0
        if alloc_path:
            alloc_moved = apply_allocation(items, store_names, alloc_path)
            alloc_src = Path(alloc_path).name
        event_src, event_moved = None, 0
        if event_path:
            event_moved = apply_event(items, store_names, event_path)
            event_src = Path(event_path).name
            n = len(store_names)
        event_stock, event_stock_added = None, 0
        if event_stock_path:
            event_stock, event_stock_added = apply_event_stock(items, store_names, event_stock_path)
        qty = [sum(i["qty"][k] for i in items) for k in range(n)]
        ret = [sum(i["ret"][k] or 0 for i in items) for k in range(n)]
        total_row = {"qty": qty + [sum(qty)], "ret": ret + [sum(ret)]}
        office_src = Path(office_path).name
        shutil.copyfile(office_path, OFFICE_XLSX_COPY)
        print(f"office: {matched} rows matched, {added} rows added (office only)")

    m = re.search(r"(\d{2}/\d{2}/\d{4})", str(as_of))
    data = {
        "title": title,
        "asOf": as_of,
        "asOfDate": m.group(1) if m else "",
        "source": Path(path).name,
        "generatedAt": datetime.now(timezone(timedelta(hours=7))).strftime("%Y-%m-%d %H:%M"),  # タイ時間
        "officeSource": office_src,
        "allocSource": alloc_src if office_path else None,
        "allocMoved": alloc_moved if office_path else 0,
        "eventSource": event_src if office_path else None,
        "eventMoved": event_moved if office_path else 0,
        "eventStock": ({"cutoff": event_stock["cutoff"], "start": event_stock["startTotal"],
                        "sold": event_stock["soldFromStart"], "remain": event_stock_added}
                       if office_path and event_stock else None),
        "eventIndex": (store_names.index(EVENT_NAME) if office_path and event_src else None),  # Store Total に含めない
        "officeIndex": 0 if office_src else None,  # Office 列 (Return なし)
        "stores": store_names,
        "items": items,
        "totalRow": total_row,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    shutil.copyfile(path, XLSX_COPY)
    total = sum(sum(i["qty"]) for i in items)
    print(f"{len(items)} rows, {len(store_names)} locations, total qty {total} -> {OUT}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/Store_Stock_092526.xlsx",
         sys.argv[2] if len(sys.argv) > 2 else None,
         sys.argv[3] if len(sys.argv) > 3 else None,
         sys.argv[4] if len(sys.argv) > 4 else None,
         sys.argv[5] if len(sys.argv) > 5 else None)
